import os
import tempfile
import requests
import chromadb
import io
import urllib3
from concurrent.futures import ThreadPoolExecutor

from bs4 import BeautifulSoup
from sentence_transformers import SentenceTransformer
from pypdf import PdfReader
from docx import Document

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

NEXTCLOUD_URL = os.getenv("NEXTCLOUD_URL")
NEXTCLOUD_USER = os.getenv("NEXTCLOUD_USER")
NEXTCLOUD_PASS = os.getenv("NEXTCLOUD_PASS")

CHROMA_PATH = os.getenv("CHROMA_PATH", "/chroma_db")

client = chromadb.PersistentClient(path=CHROMA_PATH)
collection = client.get_or_create_collection("nextcloud_docs")

embedding_model = SentenceTransformer("all-MiniLM-L6-v2")

WEBDAV_URL = f"{NEXTCLOUD_URL}/remote.php/dav/files/{NEXTCLOUD_USER}/"

def get_nextcloud_files():
    response = requests.request(
        "PROPFIND", WEBDAV_URL, auth=(NEXTCLOUD_USER, NEXTCLOUD_PASS),
        headers={"Depth": "infinity"}, verify=False
    )
    soup = BeautifulSoup(response.text, "xml")
    return [
        href.text for href in soup.find_all("d:href")
        if href.text.endswith((".txt", ".md", ".pdf", ".docx", "xlsx"))
    ]

def download_file(file_path):
    response = requests.get(
        f"{NEXTCLOUD_URL}{file_path}",
        auth=(NEXTCLOUD_USER, NEXTCLOUD_PASS),
        verify=False
    )
    return response.content

def extract_text(file_path, content):
    try:
        if file_path.endswith((".txt", ".md")):
            return content.decode("utf-8", errors="ignore")
        elif file_path.endswith(".pdf"):
            with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
                tmp.write(content)
                tmp.flush()
                return "\n".join(
                    page.extract_text() or ""
                    for page in PdfReader(tmp.name).pages
                )
        elif file_path.endswith(".docx"):
            with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
                tmp.write(content)
                tmp.flush()
                return "\n".join(p.text for p in Document(tmp.name).paragraphs)
	elif file_path.endswith(".xlsx"):
    	    import openpyxl
    	        wb = openpyxl.load_workbook(io.BytesIO(content))
    		return "\n".join(str(cell.value) for sheet in wb for row in sheet for cell in row if cell.value)
    except Exception as e:
        print(f"Extraction error for {file_path}: {e}")
    return ""

def split_text(text, chunk_size=2000):
    text = text.strip()
    return [text[i:i + chunk_size] for i in range(0, len(text), chunk_size) if text[i:i + chunk_size].strip()]

def clear_existing_entries(file_path):
    try:
        existing = collection.get(where={"path": file_path})
        if existing["ids"]:
            collection.delete(ids=existing["ids"])
            print(f"Removed old entries for {file_path}")
    except Exception as e:
        print(f"Cleanup failed for {file_path}: {e}")

def process_file(file_path):
    try:
        print(f"Indexing {file_path}")
        clear_existing_entries(file_path)
        text = extract_text(file_path, download_file(file_path))
        if not text.strip():
            print(f"No text extracted: {file_path}")
            return
        chunks = split_text(text)
        if not chunks:
            print(f"No chunks created: {file_path}")
            return
        for idx, chunk in enumerate(chunks):
            collection.add(
                ids=[f"{hash(file_path)}_{idx}"],
                documents=[chunk],
                embeddings=[embedding_model.encode(chunk).tolist()],
                metadatas=[{
                    "owner": NEXTCLOUD_USER,
                    "path": file_path,
                    "filename": os.path.basename(file_path),
                    "chunk": idx,
                    "source": "nextcloud"
                }]
            )
        print(f"Indexed {file_path} ({len(chunks)} chunks)")
    except Exception as e:
        print(f"Error indexing {file_path}: {e}")

def index_documents():
    files = get_nextcloud_files()
    print(f"Found {len(files)} files")
    with ThreadPoolExecutor(max_workers=4) as executor:
        executor.map(process_file, files)

if __name__ == "__main__":
    index_documents()
